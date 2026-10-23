from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook

from easy_tools.core.question_bank import QuestionAnswerService, load_question_bank
from easy_tools.config.schema import load_defaults, normalize, validate


def test_excel_question_bank_and_original_answer(tmp_path):
    path = tmp_path / "bank.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "基础题"
    ws.append(["问题", "答案"])
    ws.append(["牛顿第二定律是什么？", "力等于质量乘以加速度（F=ma）。\n这是原答案第二行。"])
    wb.save(path)

    entries = load_question_bank(path)
    assert entries[0].row == "基础题!2"
    assert "原答案第二行" in entries[0].answer

    class Store:
        def section(self, key, default=None):
            if key == "question_banks":
                return [{"id": "b1", "name": "物理题库", "path": str(path), "enabled": True}]
            return default

    result = QuestionAnswerService(Store()).answer("请问牛顿第二定律是什么？")
    assert result["matched"] is True
    assert result["answer"] == entries[0].answer
    assert result["generated"] is False
    assert result["matches"][0]["source"] == "物理题库"


def test_legacy_doc_question_bank_is_rejected_with_conversion_hint(tmp_path):
    import pytest
    path = tmp_path / "bank.doc"
    path.write_bytes(b"not-an-excel-file")
    with pytest.raises(ValueError, match="另存为 .docx"):
        load_question_bank(path)


def _write_docx(path, body):
    xml = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f'<w:body>{body}<w:sectPr/></w:body></w:document>')
    with ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", xml)


def test_word_knowledge_bank_extracts_headings_paragraphs_and_tables(tmp_path):
    path = tmp_path / "运维手册.docx"
    heading = '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>模型升级</w:t></w:r></w:p>'
    paragraph = ('<w:p><w:r><w:t>极速版下载的模型保存在用户数据目录，更新软件后会自动复用，'
                 '不需要再次下载。</w:t></w:r></w:p>')
    table = ('<w:tbl><w:tr><w:tc><w:p><w:r><w:t>更新策略</w:t></w:r></w:p></w:tc>'
             '<w:tc><w:p><w:r><w:t>模型文件保留在本机</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
    _write_docx(path, heading + paragraph + table)

    entries = load_question_bank(path)

    assert len(entries) == 1
    assert all(entry.kind == "document" for entry in entries)
    assert "模型升级" in entries[0].question
    assert "用户数据目录" in entries[0].answer
    assert "模型文件保留" in entries[0].answer
    assert "第 2 段" in entries[0].row
    assert "表格 1 第 1 行" in entries[0].row


def test_word_knowledge_bank_retrieval_returns_cited_document_evidence(tmp_path):
    path = tmp_path / "模型说明.docx"
    paragraph = ('<w:p><w:r><w:t>轻量版模型文件保存在本机用户数据目录。升级 Easy-Tools 后，'
                 '程序会自动加载已有模型文件，用户不需要重复下载。</w:t></w:r></w:p>')
    _write_docx(path, paragraph)

    class Store:
        def section(self, key, default=None):
            if key == "question_banks":
                return [{"id": "word-1", "name": "模型说明", "path": str(path), "enabled": True}]
            return default

    service = QuestionAnswerService(Store())
    result = service.retrieve("软件升级后是否要重新下载本机模型？")

    assert result
    assert result[0]["kind"] == "document"
    assert result[0]["source"] == "模型说明"
    assert "第 1 段" in result[0]["row"]
    assert "不需要重复下载" in result[0]["answer"]


def test_word_evidence_is_sent_to_model_and_not_returned_as_an_excel_answer():
    from types import SimpleNamespace
    from unittest.mock import Mock
    from easy_tools.core.app import AppService
    from easy_tools.core.question_bank import prefers_original_answer, verified_bank_answer

    evidence = [{"kind": "document", "source": "模型说明.docx", "row": "更新 · 第 3 段",
                 "question": "模型升级说明", "answer": "模型会在升级后自动复用。", "score": 0.7}]
    service = object.__new__(AppService)
    service.question_answer = SimpleNamespace(requires_document_evidence=lambda: True)
    service.store = SimpleNamespace(section=lambda *_: {
        "enabled": True, "provider": "builtin", "answer_content_only": False})
    service.builtin_model = SimpleNamespace(answer=Mock(return_value="升级后自动加载已有模型。"))
    service.ccswitch_model = SimpleNamespace(answer=Mock())

    assert service.answer_question("升级后还要下载模型吗？", evidence) == "升级后自动加载已有模型。"
    service.builtin_model.answer.assert_called_once_with("升级后还要下载模型吗？", evidence)
    assert not prefers_original_answer(evidence)
    assert verified_bank_answer(evidence) is None


def test_excel_uses_correct_answer_column_not_question_type(tmp_path):
    path = tmp_path / "choice-bank.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["题干", "题型", "选项A", "选项B", "正确答案"])
    ws.append(["下列哪个是质数？", "单选", "9", "11", "B"])
    wb.save(path)

    entry = load_question_bank(path)[0]
    assert entry.question == "下列哪个是质数？"
    assert entry.answer == "B. 11"


def test_choice_bank_without_answer_header_is_rejected(tmp_path):
    path = tmp_path / "missing-answer.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["题干", "选项A", "选项B"])
    ws.append(["2+2=?", "3", "4"])
    wb.save(path)

    import pytest
    with pytest.raises(ValueError, match="没有找到正确答案列"):
        load_question_bank(path)


def test_actual_bank_headers_prioritize_correct_answer_and_numbered_options(tmp_path):
    path = tmp_path / "numbered-options.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["题目名称", "题目类型", "题目难度", "答案1", "答案2", "答案3", "正确答案", "答案", "解析"])
    ws.append(["中国的首都是哪里？", "单选", "容易", "上海", "北京", "广州", "B", "A", "干扰解析"])
    ws.append(["哪些数是质数？", "多选", "容易", "2", "4", "5", "Ａ、Ｃ", "B", "干扰解析"])
    wb.save(path)
    entries = load_question_bank(path)
    assert entries[0].answer == "B. 北京"
    assert entries[0].answer_column == "正确答案"
    assert entries[0].answer_raw == "B"
    assert entries[1].answer == "A. 2\nC. 5"


def test_english_correct_answer_beats_generic_answer_column(tmp_path):
    path = tmp_path / "english.xlsx"
    wb = Workbook()
    wb.active.append(["question", "答案", "Correct Answer"])
    wb.active.append(["2 + 2?", "3", "4"])
    wb.save(path)
    assert load_question_bank(path)[0].answer == "4"


def test_unrecognized_question_header_never_falls_back_to_question_type(tmp_path):
    import pytest
    path = tmp_path / "unknown-header.xlsx"
    wb = Workbook()
    wb.active.append(["未知题干字段", "题目类型", "正确答案"])
    wb.active.append(["2 + 2?", "单选", "4"])
    wb.save(path)
    with pytest.raises(ValueError, match="题目列"):
        load_question_bank(path)


def test_verified_correct_answer_cannot_be_overwritten_by_model():
    from types import SimpleNamespace
    from unittest.mock import Mock
    from easy_tools.core.app import AppService
    from easy_tools.core.question_bank import format_bank_evidence

    evidence = [{"question": "中国的首都是哪里？", "answer": "B. 北京", "score": 1.0,
                 "answer_column": "正确答案", "answer_raw": "B"}]
    for provider in ("builtin", "ccswitch"):
        service = object.__new__(AppService)
        service.store = SimpleNamespace(section=lambda *args: {
            "enabled": True, "prefer_bank_answer": False, "provider": provider})
        service.builtin_model = Mock()
        service.ccswitch_model = Mock()
        assert service.answer_question("中国的首都是哪里？", evidence) == "B. 北京"
        service.builtin_model.answer.assert_not_called()
        service.ccswitch_model.answer.assert_not_called()
    text = format_bank_evidence(evidence)
    assert "答案来源列：正确答案" in text
    assert "答案列原值：B" in text


def test_config_rejects_more_than_three_enabled_banks():
    cfg = normalize(load_defaults())
    cfg["question_banks"] = [
        {"id": str(i), "name": f"题库{i}", "path": f"C:/bank{i}.xlsx", "enabled": True}
        for i in range(4)]
    assert "最多同时启用 3 个知识库" in validate(cfg)
